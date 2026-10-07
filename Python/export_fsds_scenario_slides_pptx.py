#!/usr/bin/env python3
"""Export scenario slide PNGs to PowerPoint (EN deck, ZH deck, combined)."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SLIDE_DIR = ROOT / "artifacts" / "fsds_scenario_slides"
CONFIG = ROOT / "config" / "fsds_economics_assumptions.json"

ORDER = [
    "01_cross_border_copilot",
    "02_bank_insurance_fl",
    "03_group_uplift",
    "04_agent_platform",
    "05_bpo_hybrid",
    "06_feature_store_ofs",
]


def _title_slide(prs, title: str, subtitle: str) -> None:
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(12), Inches(2.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.size = Pt(14)


def _assumptions_slide(prs, lang: str) -> None:
    cfg = {}
    if CONFIG.is_file():
        cfg = json.loads(CONFIG.read_text(encoding="utf-8"))
    u = cfg.get("uplift") or {}
    a = cfg.get("agent_unit_economics_from_demos") or {}
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    title = "Assumptions 假设参数" if lang == "both" else ("Assumptions" if lang == "en" else "假设参数")
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.3), Inches(6.8))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(22)
    p.font.bold = True
    lines = [
        f"margin_usd={u.get('margin_usd_per_conversion', 48)} · treat_cost={u.get('treatment_cost_usd', 6.5)}",
        f"SoT save ~{a.get('so_cost_reduction_pct', 33.6)}% · debate ~{a.get('debate_cost_reduction_pct', 39.3)}% (demos)",
        "Edit: config/fsds_economics_assumptions.json → re-run run_fsds_business_impact_pack.py",
        "OPEX vs CAPEX: RELEARN tickets excluded from OPEX net in uplift rollup",
    ]
    for line in lines:
        para = tf.add_paragraph()
        para.text = line
        para.font.size = Pt(13)


def _build_deck(lang: str, out_path: Path, *, include_assumptions: bool = True) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    if lang == "en":
        _title_slide(prs, "FSDS scenario one-pagers", "Data digestion · Economic justification · v2")
    elif lang == "zh":
        _title_slide(prs, "FSDS 场景一页纸", "数据消化 · 经济效益 · v2")
    else:
        _title_slide(prs, "FSDS scenarios / 场景", "Bilingual deck · v2")

    if include_assumptions:
        _assumptions_slide(prs, lang)

    suffix = lang if lang in ("en", "zh") else "en"
    for stem in ORDER:
        if lang == "both":
            for suf in ("en", "zh"):
                path = SLIDE_DIR / f"slide_{stem}_{suf}.png"
                if path.is_file():
                    slide = prs.slides.add_slide(blank)
                    slide.shapes.add_picture(
                        str(path), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height
                    )
        else:
            path = SLIDE_DIR / f"slide_{stem}_{suffix}.png"
            if not path.is_file():
                raise FileNotFoundError(path)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(
                str(path), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height
            )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print("Wrote", out_path)


def main() -> None:
    _build_deck("en", SLIDE_DIR / "fsds_scenario_deck_en.pptx")
    _build_deck("zh", SLIDE_DIR / "fsds_scenario_deck_zh.pptx")
    _build_deck("both", SLIDE_DIR / "fsds_scenario_deck_bilingual.pptx", include_assumptions=True)


if __name__ == "__main__":
    main()
