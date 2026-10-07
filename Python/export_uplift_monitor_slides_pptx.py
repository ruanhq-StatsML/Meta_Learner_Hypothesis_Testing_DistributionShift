#!/usr/bin/env python3
"""Export uplift monitor slide PNGs to PowerPoint (16:9)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
SLIDE_DIR = ROOT / "artifacts" / "uplift_monitor_slides"
OUT = ROOT / "artifacts" / "uplift_monitor_slides" / "uplift_two_batch_monitor_deck.pptx"

SLIDES = [
    "slide1_two_batch_monitor_pipeline.png",
    "slide2_detect_drperm_fsds.png",
    "slide3_adapt_priority_online.png",
]


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]  # blank

    for name in SLIDES:
        path = SLIDE_DIR / name
        if not path.is_file():
            raise FileNotFoundError(path)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
